#!/usr/bin/env python3
"""Contract fixtures for each PiFlow adaptation; publication requires a passing fixture."""
import importlib.util
import io
import json
from pathlib import Path
import sys
import tempfile

import pyarrow as pa

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[2] / 'dacp-operator-sdk/python'))
sys.path.insert(0, str(HERE / 'operators'))
sys.path.extend(str(p) for p in (HERE / 'target/excel-deps').glob('*.whl'))
from dacp_operator import OperatorContext
support_spec = importlib.util.spec_from_file_location('dacp_piflow_support', HERE / 'operators/support.py')
support = importlib.util.module_from_spec(support_spec)
sys.modules[support_spec.name] = support
support_spec.loader.exec_module(support)
sql_spec = importlib.util.spec_from_file_location('dacp_piflow_sql', HERE / 'operators/execute_sql_stop.py')
sql_module = importlib.util.module_from_spec(sql_spec)
sys.modules[sql_spec.name] = sql_module
sql_spec.loader.exec_module(sql_module)


def execute(name, params, inputs):
    spec = importlib.util.spec_from_file_location(name, HERE / 'operators' / (name + '.py'))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    events = []
    with tempfile.TemporaryDirectory() as tmp:
        directory = Path(tmp)
        entries = []
        for port, table in inputs.items():
            representation = 'records'
            if isinstance(table, bytes):
                representation = 'bytes'
                table = pa.table({'data': pa.array([table[i:i+3] for i in range(0, len(table), 3)], type=pa.binary())})
            with (directory / (port + '.arrow')).open('wb') as f:
                with pa.ipc.new_stream(f, table.schema) as writer:
                    writer.write_table(table, max_chunksize=1)
            entries.append(dict(port=port, file=port + '.arrow', representation=representation, format='arrow'))
        (directory / 'manifest.json').write_text(json.dumps(entries))
        (directory / 'params.json').write_text(json.dumps(params))
        with OperatorContext(['operator', tmp, str(directory / 'params.json')], emit=lambda h, b: events.append((h, b))) as ctx:
            module.process(ctx)
    result = {}
    representations = {}
    for header, data in events:
        port = header['port']
        if header['type'] == 'start':
            representations[port] = header['representation']
            result[port] = [] if header['representation'] == 'records' else b''
        elif header['type'] == 'data':
            if representations[port] == 'records':
                result[port].extend(pa.ipc.open_stream(data).read_all().to_pylist())
            else:
                result[port] += data
    assert events and events[-1][0]['type'] == 'end'
    return result


def verify(name):
    table = pa.table(dict(station_id=['A', 'B', 'A'], temperature=[12., None, 12.], quality=[0, 1, 0]))
    if name == 'reichstein_check':
        rows=[]
        for group in range(6):
            for i in range(40):
                wind=((i*7)%40)%20/10
                rows.append(dict(soil=group*100.+i,wind=wind,Fc=1. if wind==0 else 10.,day_status='night'))
        p=dict(soilTemperatureField='soil',frictionWindSpeedField='wind',FC='Fc',flagValue=-9999,percentFlux=.99)
        actual=execute(name,p,dict(source=pa.Table.from_pylist(rows)))['result']
        assert len(actual)==240 and all(r['Fc']==(-9999. if r['wind']==0 else 10.) for r in actual)
        try:
            execute(name,p,dict(source=pa.Table.from_pylist(rows[:10])))
        except ValueError:
            pass
        else:
            raise AssertionError('undersized estimation sample accepted')
    elif name == 'excel_read':
        import openpyxl
        wb=openpyxl.Workbook()
        ws=wb.active;ws.title='Weather Station'
        ws.append(['note']);ws.append([None,'station','temp']);ws.append([None,'东站',12.5]);ws.append([None,'西站',None])
        wb.create_sheet('Other').append(['other'])
        buffer=io.BytesIO();wb.save(buffer);wb.close()
        assert execute(name, dict(dataAddress="'Weather Station'!B2:C4"), dict(source=buffer.getvalue())) == {'result':[dict(station='东站',temp='12.5'),dict(station='西站',temp=None)]}
        assert execute(name, dict(dataAddress="'Other'!A1",header=False), dict(source=buffer.getvalue())) == {'result':[{'_0':'other'}]}
    elif name == 'excel_write':
        import openpyxl
        content = execute(name, dict(dataAddress="'Weather Station'!B2"), dict(source=table))['result']
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
        try:
            ws = wb['Weather Station']
            assert ws['B2'].value == 'station_id' and ws['C3'].value == 12.
            assert ws['C4'].value is None and ws['B5'].value == 'A'
        finally:
            wb.close()
        formula = pa.table(dict(text=['=1+1', '天气']))
        wb = openpyxl.load_workbook(io.BytesIO(execute(name, {}, dict(source=formula))['result']))
        try:
            assert wb.active['A2'].data_type=='s' and wb.active['A2'].value=='=1+1'
            assert wb.active['A3'].value=='天气'
        finally:
            wb.close()
    elif name == 'mock_data':
        rows = execute(name, dict(schema='name:String,age:Int,flag:Boolean,value:Double:true', count='1030'), dict(source=table))['result']
        assert len(rows)==1030
        assert all(len(r['name'])==10 and 0<=r['age']<100 and type(r['flag']) is bool for r in rows)
        assert execute(name, dict(schema='n:Long', count='0'), dict(source=table)) == {'result': []}
    elif name == 'select_field':
        actual = execute(name, dict(columnNames='temperature,station_id'), dict(source=table))
        assert actual == dict(result=[dict(temperature=12., station_id='A'), dict(temperature=None, station_id='B'), dict(temperature=12., station_id='A')])
        assert execute(name, dict(columnNames='station_id'), dict(source=table.slice(0, 0))) == {'result': []}
        for columns in ('missing', '', 'station_id,station_id'):
            try:
                execute(name, dict(columnNames=columns), dict(source=table))
            except ValueError:
                pass
            else:
                raise AssertionError('invalid columns accepted')
    elif name == 'drop_field':
        actual = execute(name, dict(columnNames='quality,absent'), dict(source=table))
        assert actual == dict(result=[dict(station_id='A', temperature=12.), dict(station_id='B', temperature=None), dict(station_id='A', temperature=12.)])
        assert execute(name, dict(columnNames='quality'), dict(source=table.slice(0, 0))) == {'result': []}
    elif name == 'convert_schema':
        actual = execute(name, dict(schema='station_id->station,temperature->air_c'), dict(source=table))
        assert actual == dict(result=[dict(station='A', air_c=12., quality=0), dict(station='B', air_c=None, quality=1), dict(station='A', air_c=12., quality=0)])
        try:
            execute(name, dict(schema='station_id->quality'), dict(source=table))
        except ValueError:
            pass
        else:
            raise AssertionError('duplicate output columns accepted')
    elif name == 'distinct':
        for p in ({}, dict(columnNames='station_id')):
            actual = execute(name, p, dict(source=table))
            assert actual == {'result': table.slice(0, 2).to_pylist()}
        assert execute(name, {}, dict(source=table.slice(0, 0))) == {'result': []}
    elif name == 'add_uuid_stop':
        import re
        rows = execute(name, dict(column='uuid'), dict(source=table))['result']
        ids = [row.pop('uuid') for row in rows]
        assert rows == table.to_pylist()
        assert len(set(ids)) == len(ids) and all(re.fullmatch('[0-9a-f]{32}', i) for i in ids)
    elif name == 'filter':
        assert execute(name, dict(condition='quality = 0 AND temperature >= 12'), dict(source=table)) == {'result': [table.to_pylist()[0], table.to_pylist()[2]]}
        assert execute(name, dict(condition='temperature IS NULL'), dict(source=table)) == {'result': [table.to_pylist()[1]]}
        assert execute(name, dict(condition='0'), dict(source=table)) == {'result': []}
        try:
            execute(name, dict(condition='1); DROP TABLE source; --'), dict(source=table))
        except Exception:
            pass
        else:
            raise AssertionError('multiple SQL statements accepted')
    elif name == 'fork':
        assert execute(name, {}, dict(source=table)) == {'left': table.to_pylist(), 'right': table.to_pylist()}
        assert execute(name, {}, dict(source=table.slice(0, 0))) == {'left': [], 'right': []}
    elif name == 'column_check':
        assert execute(name, {}, dict(left=table, right=table.slice(0, 0))) == {'result': table.to_pylist()}
        for right in (table.select(['quality', 'temperature', 'station_id']), table.select(['station_id'])):
            try:
                execute(name, {}, dict(left=table, right=right))
            except ValueError:
                pass
            else:
                raise AssertionError('mismatching reference accepted')
    elif name == 'consistency_check':
        right = pa.table(dict(station_id=['A', 'A', None]))
        actual = execute(name, dict(flagField='flag', correlationColumn='station_id'), dict(left=table, right=right))['result']
        assert [r['flag'] for r in actual] == [0, 0, 1, 0, 0]
        assert execute(name, dict(flagField='flag', correlationColumn='station_id'), dict(left=table.slice(0, 0), right=right)) == {'result': []}
    elif name == 'merge':
        assert execute(name, {}, dict(left=table, right=table.slice(0, 1))) == {'result': table.to_pylist() + table.slice(0, 1).to_pylist()}
        assert execute(name, {}, dict(left=table.slice(0, 0), right=table)) == {'result': table.to_pylist()}
        try:
            execute(name, {}, dict(left=table, right=table.select(['station_id'])))
        except ValueError:
            pass
        else:
            raise AssertionError('incompatible merge accepted')
    elif name == 'subtract':
        assert execute(name, {}, dict(left=table, right=table.slice(1, 1))) == {'result': table.slice(0, 1).to_pylist()}
        assert execute(name, {}, dict(left=table, right=table.slice(0, 0))) == {'result': table.slice(0, 2).to_pylist()}
        assert execute(name, {}, dict(left=table, right=table)) == {'result': []}
    elif name == 'join':
        right = pa.table(dict(station_id=['A', 'C'], reference=[10., 11.]))
        for mode, count in [('inner', 2), ('left', 3), ('right', 3), ('full_outer', 4)]:
            rows = execute(name, dict(joinMode=mode, correlationColumn='station_id'), dict(left=table, right=right))['result']
            assert len(rows) == count
            assert all(row['reference'] == 10. for row in rows if row['station_id'] == 'A')
        assert execute(name, dict(joinMode='inner', correlationColumn='station_id'), dict(left=table, right=right.slice(0, 0))) == {'result': []}
    elif name == 'route':
        actual = execute(name, dict(condition='temperature >= 12'), dict(source=table))
        assert actual == {'result': table.to_pylist(), 'matched': [table.to_pylist()[0], table.to_pylist()[2]], 'unmatched': [table.to_pylist()[1]]}
    elif name == 'execute_sql_stop':
        actual = execute(name, dict(viewName='observations', sql='SELECT station_id,COUNT(*) AS samples,AVG(temperature) AS mean_c FROM observations GROUP BY station_id ORDER BY station_id'), dict(source=table))
        assert actual == {'result': [dict(station_id='A', samples=2, mean_c=12.), dict(station_id='B', samples=1, mean_c=None)]}
        assert execute(name, dict(viewName='source', sql='SELECT * FROM source WHERE 0'), dict(source=table)) == {'result': []}
        for sql in ('DROP TABLE source', 'ATTACH DATABASE "/tmp/piflow.db" AS extra'):
            try:
                execute(name, dict(viewName='source', sql=sql), dict(source=table))
            except Exception:
                pass
            else:
                raise AssertionError('write SQL accepted')
    elif name in ('csv_save', 'custom_view'):
        actual = execute(name, {}, dict(source=table))
        assert actual == {'result': b'station_id,temperature,quality\nA,12.0,0\nB,,1\nA,12.0,0\n'}
        special = pa.table({'text': ['a,b', 'line\nnext', '"quoted"', '天气']})
        import csv
        text = execute(name, {}, dict(source=special))['result'].decode()
        assert list(csv.reader(io.StringIO(text))) == [['text'], ['a,b'], ['line\nnext'], ['"quoted"'], ['天气']]
    elif name == 'csv_parser':
        raw = 'station,temp,note\n东站,12.5,"two\nlines"\n西站,,"a,b"\n'.encode()
        assert execute(name, {}, dict(source=raw)) == {'result': [dict(station='东站', temp='12.5', note='two\nlines'), dict(station='西站', temp='', note='a,b')]}
        assert execute(name, dict(header=False, schema='station,temp'), dict(source=b'A,12\nB,13\n')) == {'result': [dict(station='A', temp='12'), dict(station='B', temp='13')]}
        assert execute(name, {}, dict(source=b'station,temp\n')) == {'result': []}
        try:
            execute(name, {}, dict(source=b'a,b\n1,2,3\n'))
        except ValueError:
            pass
        else:
            raise AssertionError('wrong CSV width accepted')
    elif name == 'csv_string_parser':
        assert execute(name, dict(schema='station,temp', delimiter=';'), dict(source=b'A;12\nB;13\n')) == {'result': [dict(station='A', temp='12'), dict(station='B', temp='13')]}
        assert execute(name, {}, dict(source=b'A,12\n')) == {'result': [{'_0': 'A', '_1': '12'}]}
    elif name == 'table_show':
        assert execute(name, dict(showField='station_id, temperature'), dict(source=table)) == {'result': [{'station_id': r['station_id'], 'temperature': r['temperature']} for r in table.to_pylist()]}
    elif name == 'scatter_plot_chart':
        assert execute(name, dict(legend='station_id', customizedProperties={'quality':1, 'temperature':0}), dict(source=table)) == {'result': [dict(station_id='A', temperature=12., quality=0), dict(station_id='A', temperature=12., quality=0), dict(station_id='B', temperature=None, quality=1)]}
    elif name == 'pie_chart':
        assert execute(name, dict(dimension='station_id', indicator='temperature', indicatorOption='AVG'), dict(source=table)) == {'result': [dict(station_id='A', temperature_AVG=12.), dict(station_id='B', temperature_AVG=None)]}
    elif name in ('histogram', 'line_chart'):
        assert execute(name, dict(abscissa='station_id', customizedProperties={'temperature':'AVG', 'quality':'SUM'}), dict(source=table)) == {'result': [dict(station_id='A', temperature_AVG=12., quality_SUM=0), dict(station_id='B', temperature_AVG=None, quality_SUM=1)]}
    else:
        raise ValueError('no verification fixture: ' + name)
    print('PASS ' + name, flush=True)


if __name__ == '__main__':
    for name in sys.argv[1:]:
        verify(name)
