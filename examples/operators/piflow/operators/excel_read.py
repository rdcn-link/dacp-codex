"""Read a selected XLSX worksheet/range from an input byte stream as string records."""
import datetime
import io
import re
import zipfile
import pyarrow as pa
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple
from dacp_piflow_support import params, MAX_BYTES


def process(ctx):
    p = params(ctx, optional=['header','dataAddress'])
    header, address=p.get('header',True),p.get('dataAddress','A1')
    if type(header) is not bool or not isinstance(address,str):
        raise ValueError('header must be boolean; dataAddress must be string')
    match=re.fullmatch(r"(?:(?:'((?:[^']|'')+)'|([^!]+))!)?([A-Za-z]+[1-9][0-9]*(?::[A-Za-z]+[1-9][0-9]*)?)",address)
    if not match:
        raise ValueError('dataAddress must be [Sheet!]A1[:F35]')
    buffer=io.BytesIO()
    for chunk in ctx.input('source').bytes():
        if buffer.tell()+len(chunk)>MAX_BYTES:
            raise ValueError('XLSX input exceeds 128 MiB')
        buffer.write(chunk)
    buffer.seek(0)
    with zipfile.ZipFile(buffer) as archive:
        if sum(x.file_size for x in archive.infolist())>MAX_BYTES:
            raise ValueError('expanded XLSX exceeds 128 MiB')
    buffer.seek(0)
    wb=load_workbook(buffer,read_only=True,data_only=True,keep_links=False)
    try:
        sheet=(match[1] or match[2] or '').replace("''", "'")
        ws=wb[sheet] if sheet else wb.worksheets[0]
        if ':' in match[3]:
            c1,r1,c2,r2=range_boundaries(match[3])
        else:
            r1,c1=coordinate_to_tuple(match[3]);r2,c2=ws.max_row,ws.max_column
        if not r2 or not c2 or r2<r1 or c2<c1 or (r2-r1+1)*(c2-c1+1)>1000000:
            raise ValueError('selected XLSX range must contain at most one million cells')
        rows=iter(ws.iter_rows(min_row=r1,max_row=r2,min_col=c1,max_col=c2,values_only=True))
        first=next(rows,None)
        if header:
            if first is None or any(v is None for v in first):
                raise ValueError('nonempty XLSX header required')
            names=[str(v) for v in first]
        else:
            import itertools
            names=['_'+str(i) for i in range(c2-c1+1)]
            rows=itertools.chain([] if first is None else [first],rows)
        if any(not n for n in names) or len(set(names))!=len(names):
            raise ValueError('unique nonempty XLSX column names required')
        schema=pa.schema([(n,pa.string()) for n in names])
        out=ctx.records_output('result',schema)
        batch=[]
        for row in rows:
            def text(v):
                if isinstance(v,(datetime.date,datetime.datetime,datetime.time)):
                    return v.isoformat()
                if isinstance(v,bool):
                    return 'true' if v else 'false'
                return None if v is None else str(v)
            batch.append(dict(zip(names,map(text,row))))
            if len(batch)==128:
                out.write(pa.RecordBatch.from_pylist(batch,schema=schema));batch.clear()
        if batch:
            out.write(pa.RecordBatch.from_pylist(batch,schema=schema))
    finally:
        wb.close()


if __name__ == '__main__':
    from dacp_operator import OperatorContext
    with OperatorContext() as context:
        process(context)
